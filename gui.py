"""
gui.py - 货运台账解析器 GUI
v2.1: 自定义保存路径 + 司机类型管理
对标 MiniMax Code 输出格式
"""
import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog, simpledialog
import os, re
from datetime import datetime

from robust_parser import parse_driver_message, format_report_text, RowData
from config import load_config, save_config

# ========== 颜色 ==========
BG_COLOR = "#f0f0f0"
BTN_BG = "#4a90d9"
BTN_FG = "white"

# ========== 默认保存路径 ==========
DEFAULT_SAVE_DIR = r"C:\Users\ADMIN\Desktop\每天司机数据"


class FreightLedgerApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("司机发货登记 v2.1")
        self.root.geometry("1100x780")
        self.root.configure(bg=BG_COLOR)

        self.cfg = load_config()
        # 保存路径（从配置读取）
        self.save_dir = self.cfg.get("save_path", DEFAULT_SAVE_DIR)
        self._current_result = None

        self._build_menu()
        self._build_ui()

    def _build_menu(self):
        menu_bar = tk.Menu(self.root)
        self.root.config(menu=menu_bar)

        # 配置菜单
        config_menu = tk.Menu(menu_bar, tearoff=0)
        config_menu.add_command(label="📂 设置保存路径...", command=self._set_save_path)
        config_menu.add_command(label="🚚 配置司机类别...", command=self._open_driver_config)
        config_menu.add_separator()
        config_menu.add_command(label="📖 使用说明", command=self._show_help)
        menu_bar.add_cascade(label="⚙️ 配置", menu=config_menu)

    def _build_ui(self):
        main_frame = ttk.Frame(self.root, padding=10)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # ===== 保存路径显示 =====
        path_frame = ttk.Frame(main_frame)
        path_frame.pack(fill=tk.X, pady=(0, 4))
        ttk.Label(path_frame, text="📂 保存到:", font=("微软雅黑", 9)).pack(side=tk.LEFT)
        self.path_label = ttk.Label(path_frame, text=self.save_dir,
            font=("微软雅黑", 9), foreground="#666666")
        self.path_label.pack(side=tk.LEFT, padx=4)
        ttk.Button(path_frame, text="更改", command=self._set_save_path,
            width=6).pack(side=tk.LEFT, padx=4)

        # ===== 输入区 =====
        input_frame = ttk.LabelFrame(main_frame, text="📥 粘贴司机发货消息（支持多司机一起黏贴）", padding=8)
        input_frame.pack(fill=tk.BOTH, pady=(0, 6))

        self.text_input = scrolledtext.ScrolledText(
            input_frame, height=10, wrap=tk.WORD,
            font=("Consolas", 11), bg="white"
        )
        self.text_input.pack(fill=tk.BOTH, expand=True, pady=(0, 6))

        tool_row = ttk.Frame(input_frame)
        tool_row.pack(fill=tk.X)

        ttk.Label(tool_row, text="年份:").pack(side=tk.LEFT, padx=(0, 4))
        self.year_var = tk.StringVar(value=self.cfg.get("default_year", "2026"))
        ttk.Entry(tool_row, textvariable=self.year_var, width=6).pack(side=tk.LEFT, padx=(0, 12))

        ttk.Button(tool_row, text="🧹 清空", command=self._clear).pack(side=tk.LEFT, padx=4)
        ttk.Button(tool_row, text="📋 粘贴示例", command=self._paste_example).pack(side=tk.LEFT, padx=4)

        self.parse_btn = tk.Button(
            tool_row, text="🚀 开始解析",
            command=self._do_parse,
            bg=BTN_BG, fg=BTN_FG, font=("微软雅黑", 11, "bold"),
            relief=tk.RAISED, borderwidth=2, padx=20, cursor="hand2"
        )
        self.parse_btn.pack(side=tk.RIGHT, padx=4)

        # ===== 状态 =====
        self.status_label = ttk.Label(main_frame, text="就绪", font=("微软雅黑", 10))
        self.status_label.pack(fill=tk.X, pady=2)

        # ===== 输出预览区 =====
        output_frame = ttk.LabelFrame(main_frame, text="📊 解析结果预览", padding=8)
        output_frame.pack(fill=tk.BOTH, expand=True)

        btn_row = ttk.Frame(output_frame)
        btn_row.pack(fill=tk.X, pady=(0, 6))

        ttk.Button(btn_row, text="💾 保存 TXT + XLSX", command=self._save_output).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_row, text="📋 复制报告", command=self._copy_report).pack(side=tk.LEFT, padx=4)

        table_frame = ttk.Frame(output_frame)
        table_frame.pack(fill=tk.BOTH, expand=True)

        columns = RowData.header()
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings",
                                  height=12, selectmode="extended")

        col_widths = [80, 110, 80, 100, 100, 60, 120, 70, 60, 80, 200]
        for col, w in zip(columns, col_widths):
            self.tree.heading(col, text=col, anchor=tk.W)
            self.tree.column(col, width=w, minwidth=50, anchor=tk.W)

        vsb = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self.count_label = ttk.Label(output_frame, text="共 0 行")
        self.count_label.pack(anchor=tk.W, pady=(4, 0))

        self.root.bind("<Control-Return>", lambda e: self._do_parse())

    # ===== 保存路径设置 =====

    def _set_save_path(self):
        path = filedialog.askdirectory(
            title="选择保存目录",
            initialdir=self.save_dir
        )
        if path:
            self.save_dir = path
            self.path_label.config(text=path)
            # 保存到配置
            self.cfg["save_path"] = path
            save_config(self.cfg)
            self.status_label.config(text=f"📂 保存路径已更改为: {path}")

    # ===== 司机配置 =====

    def _open_driver_config(self):
        DriverConfigDialog(self.root, self.cfg, self._on_driver_config_saved)

    def _on_driver_config_saved(self, new_cfg: dict):
        self.cfg = new_cfg
        save_config(new_cfg)
        self.status_label.config(text="⚙️ 司机配置已更新")

    # ===== 解析逻辑 =====

    def _split_driver_blocks(self, text: str, all_drivers: list) -> list:
        """按司机分割原始文本"""
        blocks, buf = [], ""
        for line in text.split("\n"):
            s = line.strip()
            is_new = False
            for d in sorted(all_drivers, key=len, reverse=True):
                if s.startswith(d) and re.search(r'\d+:\d+', s):
                    if buf.strip(): blocks.append(buf.strip())
                    buf = s + "\n"; is_new = True; break
            if not is_new:
                buf += (s + "\n") if s else "\n"
        if buf.strip(): blocks.append(buf.strip())
        return blocks

    def _do_parse(self):
        raw = self.text_input.get("1.0", tk.END).strip()
        if not raw:
            messagebox.showwarning("提示", "请先粘贴司机原始消息")
            return

        year = self.year_var.get().strip()
        if not year.isdigit() or len(year) != 4:
            year = self.cfg.get("default_year", "2026")

        try:
            # 从配置读取司机类型、跟车等
            driver_types = self.cfg.get("drivers")
            riders = self.cfg.get("riders")
            special_places = self.cfg.get("special_places")

            # 按司机分割，逐个解析（防止跟车跨司机污染）
            all_drivers = list(driver_types.keys()) if driver_types else []
            blocks = self._split_driver_blocks(raw, all_drivers)

            all_rows = []
            global_seq = 0
            for block in blocks:
                result = parse_driver_message(
                    block, year=year,
                    driver_types=driver_types,
                    riders=riders,
                    special_places=special_places,
                )
                block_rows = result["rows"]
                # 流水号全局连续
                for r in block_rows:
                    global_seq += 1
                    r.流水号 = str(global_seq)
                all_rows.extend(block_rows)

            rows = all_rows

            if not rows:
                self.status_label.config(text="⚠️ 解析完成，但未生成任何行（请检查输入格式）")
                return

            self._current_result = rows
            self._refresh_table(rows)

            # 提取日期和司机信息
            drivers_set = set(r.姓名 for r in rows)
            date_str = rows[0].日期 if rows else "?"
            self.status_label.config(
                text=f"✅ 解析完成 — 共 {len(rows)} 行, {len(drivers_set)}个司机, 日期: {date_str}"
            )
        except Exception as e:
            messagebox.showerror("解析失败", f"错误：{str(e)}")
            self.status_label.config(text=f"❌ 出错: {str(e)}")

    def _save_output(self):
        if not self._current_result:
            messagebox.showwarning("提示", "还没有解析结果，请先粘贴并解析")
            return

        rows = self._current_result
        date_str = rows[0].日期 if rows else ""
        driver = rows[0].姓名 if rows else ""

        try:
            # 从日期创建文件夹
            date_match = re.match(r'(\d{4})/(\d{1,2})/(\d{1,2})', date_str)
            if date_match:
                y, m, d = date_match.groups()
                date_obj = f"{y}-{int(m):02d}-{int(d):02d}"
            else:
                date_obj = datetime.now().strftime("%Y-%m-%d")

            day_dir = os.path.join(self.save_dir, date_obj)
            os.makedirs(day_dir, exist_ok=True)

            day_suffix = date_obj[-5:].replace("-", "")

            # TXT 文件
            txt_path = os.path.join(day_dir, f"freight_output_{day_suffix}.txt")
            report_text = format_report_text(rows, date_str)
            with open(txt_path, "w", encoding="utf-8") as f:
                f.write(report_text)

            # XLSX 文件
            xlsx_path = os.path.join(day_dir, f"freight_{day_suffix}.xlsx")
            self._write_xlsx(rows, xlsx_path)

            self.status_label.config(
                text=f"✅ 已保存: {os.path.basename(txt_path)}, {os.path.basename(xlsx_path)}"
            )
            messagebox.showinfo("保存成功",
                f"📁 {day_dir}\n\n📄 {os.path.basename(txt_path)}\n📊 {os.path.basename(xlsx_path)}")
        except Exception as e:
            messagebox.showerror("保存失败", str(e))

    def _write_xlsx(self, rows: list, path: str):
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "台账"

        hdr = RowData.header()
        hf = Font(bold=True, size=11, color="FFFFFF")
        hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'), bottom=Side(style='thin'))

        for col, h in enumerate(hdr, 1):
            c = ws.cell(1, col, h)
            c.font = hf; c.fill = hfill
            c.alignment = Alignment(horizontal='center')
            c.border = thin

        for i, r in enumerate(rows, 2):
            for col, v in enumerate(r.to_list(), 1):
                ws.cell(i, col, v).border = thin

        for i, w in enumerate([10, 14, 10, 12, 12, 6, 14, 8, 6, 10, 30], 1):
            ws.column_dimensions[chr(64 + i)].width = w

        wb.save(path)

    def _copy_report(self):
        if not self._current_result:
            messagebox.showwarning("提示", "还没有解析结果")
            return
        rows = self._current_result
        date_str = rows[0].日期 if rows else ""
        report = format_report_text(rows, date_str)
        self.root.clipboard_clear()
        self.root.clipboard_append(report)
        self.status_label.config(text="📋 报告已复制到剪贴板")

    def _refresh_table(self, rows: list):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for i, row in enumerate(rows, 1):
            tag = "even" if i % 2 == 0 else "odd"
            self.tree.insert("", tk.END, values=row.to_list(), tags=(tag,))
        self.tree.tag_configure("odd", background="#ffffff")
        self.tree.tag_configure("even", background="#f4f8fc")
        self.count_label.config(text=f"共 {len(rows)} 行")

    def _clear(self):
        self.text_input.delete("1.0", tk.END)
        self._current_result = None
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.count_label.config(text="共 0 行")
        self.status_label.config(text="已清空")

    def _paste_example(self):
        self.text_input.delete("1.0", tk.END)
        self.text_input.insert("1.0", """宋江鸿 7/10 18:43:27
7月10号
姓名 宋江鸿
第一车 豪昇去4小回3小
第二车 汇兴去3小回5小
第三车 晶科达去4小回1小
第四车 晶科达去4小
第五车 凯旋去4小
第六车 煜达去1大3小回2小（中空）
第七车 豪昇去3小回3小""")

    def _show_help(self):
        messagebox.showinfo("使用说明",
            "📖 司机发货登记 — 使用说明\n\n"
            "【数据来源】\n"
            "司机通过企业微信发送送货消息，格式如下：\n"
            "\n"
            "X月X日\n"
            "姓名 （司机姓名）（跟车司机姓名）\n"
            "第一车（客户）去X大X小回X大X小\n"
            "第二车（客户）去X大X小回X大X小\n"
            "第三车（客户）去X大X小回X大X小（单片/中空）\n"
            "...\n"
            "\n"
            "实际例子：\n"
            "7月10号\n"
            "姓名 宋江鸿\n"
            "第一车 豪昇去4小回3小\n"
            "第二车 汇兴去3小回5小\n\n"
            "【操作步骤】\n"
            "1. 从企业微信司机群复制原始消息\n"
            "2. 粘贴到上方输入框（支持多司机一起贴）\n"
            "3. 点击「开始解析」或按 Ctrl+Enter\n"
            "4. 核对结果，点击「保存 TXT + XLSX」\n\n"
            "【输出格式】\n"
            "11列规范表格：\n"
            "年月 | 日期 | 姓名 | 类型 | 客户 | 数量 | 拼车 | 车号 | 流水号 | 跟车 | 原始数据\n\n"
            "【配置】\n"
            "⚙️ 配置 → 设置保存路径：更改文件保存位置\n"
            "⚙️ 配置 → 配置司机类别：添加/删除司机，修改货物类型\n\n"
            "【注意】\n"
            "原始数据列原样保留，不修改，用于核对")


# ==================== 司机配置对话框 ====================

class DriverConfigDialog:
    def __init__(self, parent, cfg: dict, callback):
        self.callback = callback
        self.cfg = dict(cfg)
        self.drivers = dict(cfg.get("drivers", {}))

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("🚚 配置司机类别")
        self.dialog.geometry("500x450")
        self.dialog.resizable(True, True)
        self.dialog.transient(parent)
        self.dialog.grab_set()

        main = ttk.Frame(self.dialog, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        # 列表
        list_frame = ttk.LabelFrame(main, text="司机列表", padding=8)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 8))

        columns = ("姓名", "类型")
        self.tree = ttk.Treeview(list_frame, columns=columns, show="headings", height=10)
        self.tree.heading("姓名", text="司机姓名")
        self.tree.heading("类型", text="货物类型")
        self.tree.column("姓名", width=200)
        self.tree.column("类型", width=150)

        vsb = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        self._refresh_list()

        # 操作按钮
        btn_frame = ttk.Frame(main)
        btn_frame.pack(fill=tk.X, pady=4)

        ttk.Label(btn_frame, text="姓名:").pack(side=tk.LEFT)
        self.name_entry = ttk.Entry(btn_frame, width=12)
        self.name_entry.pack(side=tk.LEFT, padx=4)

        self.type_combo = ttk.Combobox(btn_frame, values=["单片货", "中空货"],
                                        width=8, state="readonly")
        self.type_combo.set("单片货")
        self.type_combo.pack(side=tk.LEFT, padx=4)

        ttk.Button(btn_frame, text="➕ 添加", command=self._add_driver).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="❌ 删除选中", command=self._del_driver).pack(side=tk.LEFT, padx=2)

        # 底部按钮
        bottom = ttk.Frame(main)
        bottom.pack(fill=tk.X)
        ttk.Button(bottom, text="✅ 保存", command=self._save).pack(side=tk.RIGHT, padx=4)
        ttk.Button(bottom, text="取消", command=self.dialog.destroy).pack(side=tk.RIGHT, padx=4)

    def _refresh_list(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for name in sorted(self.drivers.keys()):
            dtype = self.drivers[name]
            self.tree.insert("", tk.END, values=(name, dtype))

    def _add_driver(self):
        name = self.name_entry.get().strip()
        if not name:
            messagebox.showwarning("提示", "请输入司机姓名", parent=self.dialog)
            return
        dtype = self.type_combo.get()
        if name in self.drivers:
            messagebox.showinfo("提示", f"司机「{name}」已存在，将覆盖旧类型", parent=self.dialog)
        self.drivers[name] = dtype
        self._refresh_list()
        self.name_entry.delete(0, tk.END)

    def _del_driver(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("提示", "请先选择要删除的司机", parent=self.dialog)
            return
        item = self.tree.item(sel[0])
        name = item["values"][0]
        if messagebox.askyesno("确认删除", f"确定删除司机「{name}」？", parent=self.dialog):
            if name in self.drivers:
                del self.drivers[name]
            self._refresh_list()

    def _save(self):
        if not self.drivers:
            messagebox.showwarning("提示", "至少需要一个司机", parent=self.dialog)
            return
        self.cfg["drivers"] = self.drivers
        self.callback(self.cfg)
        self.dialog.destroy()


def main():
    root = tk.Tk()
    app = FreightLedgerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
